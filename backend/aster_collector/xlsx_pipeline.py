from collections.abc import Iterable
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


def normalize_headers(values: list[Any]) -> list[str]:
    """Return non-empty, unique column names for pandas DataFrames."""

    seen: dict[str, int] = {}
    headers: list[str] = []
    for index, value in enumerate(values, start=1):
        base = str(value or "").strip() or f"__empty_{index}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        headers.append(base if count == 0 else f"{base}_{count + 1}")
    return headers


def iter_xlsx_frames(path: str | Path, sheet_name: str | None = None, chunk_size: int = 5000) -> Iterable[pd.DataFrame]:
    """Read an XLSX file in small DataFrame chunks.

    This avoids the memory spike caused by pandas.read_excel on large workbooks.
    The first row is treated as the header.
    """

    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = normalize_headers(list(next(rows)))
        batch: list[tuple[Any, ...]] = []

        for row in rows:
            batch.append(row)
            if len(batch) >= chunk_size:
                yield pd.DataFrame(batch, columns=headers)
                batch = []

        if batch:
            yield pd.DataFrame(batch, columns=headers)
    finally:
        workbook.close()


def summarize_xlsx(path: str | Path, sheet_name: str | None = None, sample_rows: int = 5) -> dict[str, Any]:
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = normalize_headers(list(next(rows)))
        sample = []
        for _, row in zip(range(sample_rows), rows):
            sample.append(dict(zip(headers, row)))

        return {
            "file": str(path),
            "sheet": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "headers": headers,
            "sample": sample,
        }
    finally:
        workbook.close()
