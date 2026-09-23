from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT_DIR = ROOT / "Planilhas"
DEFAULT_OUTPUT_JSON = ROOT / "docs" / "evidence" / "local_spreadsheet_sources_summary.json"
DEFAULT_OUTPUT_MD = ROOT / "docs" / "evidence" / "local_spreadsheet_sources_summary.md"

HEADER_ROWS = {
    "Gestão da Produção.xlsx": {
        "Gestao da ProdV6": 1,
    },
    "Margem_GABR_20260916.xlsx": {
        "Resumo (2)": 5,
        "Resumo": 5,
        "Desvios": 5,
        "BD": 2,
        "BD_Meta": 1,
        "TD_Meta": 1,
        "Apoio": 1,
        "Tabela de Preço": 2,
    },
}


def resolve_under_root(path: str | Path) -> Path:
    candidate = Path(path)
    if not candidate.is_absolute():
        candidate = ROOT / candidate
    resolved = candidate.resolve()
    if not (resolved == ROOT or ROOT in resolved.parents):
        raise ValueError(f"Caminho fora do workspace: {resolved}")
    return resolved


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " / ").strip()


def inspect_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    configured_headers = HEADER_ROWS.get(path.name, {})
    sheets = []
    for worksheet in workbook.worksheets:
        header_row = configured_headers.get(worksheet.title, 1)
        values = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        headers = [clean_cell(value) for value in values]
        sheets.append(
            {
                "name": worksheet.title,
                "rows": worksheet.max_row,
                "cols": worksheet.max_column,
                "header_row": header_row,
                "headers": headers,
            }
        )
    return {
        "path": str(path.relative_to(ROOT)),
        "size_bytes": path.stat().st_size,
        "modified_time_local": path.stat().st_mtime,
        "sheets": sheets,
    }


def write_markdown(summary: dict[str, Any], output: Path) -> None:
    lines = [
        "# Fontes Locais de Planilhas",
        "",
        "Resumo das planilhas em `Planilhas/` usadas para complementar lacunas do Aster.",
        "",
    ]
    for key, item in summary.items():
        lines.extend(
            [
                f"## {key}",
                "",
                f"- Arquivo: `{item['path']}`",
                f"- Tamanho: `{item['size_bytes']}` bytes",
                "",
                "| Aba | Linhas | Colunas | Linha cabecalho | Principais campos |",
                "|---|---:|---:|---:|---|",
            ]
        )
        for sheet in item["sheets"]:
            headers = [header for header in sheet["headers"] if header]
            preview = ", ".join(headers[:20])
            if len(headers) > 20:
                preview += ", ..."
            lines.append(
                "| {name} | {rows} | {cols} | {header_row} | {preview} |".format(
                    name=sheet["name"],
                    rows=sheet["rows"],
                    cols=sheet["cols"],
                    header_row=sheet["header_row"],
                    preview=preview.replace("|", "/"),
                )
            )
        lines.append("")
    output.write_text("\n".join(lines), encoding="utf-8")


def inspect_local_spreadsheet_sources(
    *,
    input_dir: str | Path = DEFAULT_INPUT_DIR,
    output_json: str | Path = DEFAULT_OUTPUT_JSON,
    output_md: str | Path = DEFAULT_OUTPUT_MD,
) -> dict[str, Any]:
    source_dir = resolve_under_root(input_dir)
    json_path = resolve_under_root(output_json)
    md_path = resolve_under_root(output_md)

    if not source_dir.exists():
        result = {
            "files": 0,
            "input_dir": str(source_dir.relative_to(ROOT)),
            "summary": {},
            "warning": "Diretorio de planilhas nao encontrado.",
        }
    else:
        files = sorted(source_dir.glob("*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
        summary = {path.stem: inspect_workbook(path) for path in files}
        result = {
            "files": len(files),
            "input_dir": str(source_dir.relative_to(ROOT)),
            "summary": summary,
            "warning": None,
        }

    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(result["summary"], ensure_ascii=False, indent=2), encoding="utf-8")
    write_markdown(result["summary"], md_path)

    return {
        **result,
        "output_json": str(json_path.relative_to(ROOT)),
        "output_md": str(md_path.relative_to(ROOT)),
    }
