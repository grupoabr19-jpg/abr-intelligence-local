from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from psycopg.types.json import Json

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from backend.archive_storage import (  # noqa: E402
    DRIVE_FOLDER_MIME,
    archive_config,
    download_drive_file,
    drive_access_token,
    export_drive_file,
    file_sha256,
    list_drive_children,
    slug,
)
from backend.aster_collector.local_spreadsheets import HEADER_ROWS  # noqa: E402
from tools.apply_migrations import connect_core_database, load_env  # noqa: E402


GOOGLE_SHEETS_MIME = "application/vnd.google-apps.spreadsheet"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME_TYPES = {"text/csv", "application/csv", "text/plain"}
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".csv"}
DEFAULT_SCAN_DEPTH = int(os.environ.get("DRIVE_SPREADSHEET_SCAN_DEPTH", "1") or "1")
GENERATED_ARCHIVE_SUFFIXES = (
    "_clean.csv",
    "_raw.jsonl",
    "_clean.parquet",
    "_manifest.json",
    ".csv.gz",
    ".jsonl",
    ".parquet",
)
DEFAULT_SOURCE_FOLDER_ID = "1e-ZFOdh5PpIOjX1xFU87uhIKHveL5EpU"
DEFAULT_SOURCE_FOLDER_URL = f"https://drive.google.com/drive/folders/{DEFAULT_SOURCE_FOLDER_ID}"
BUSINESS_SPREADSHEET_KEYWORDS = (
    "margem",
    "gabr",
    "gest",
    "produ",
    "ranking",
    "estoque",
    "envelhecimento",
    "env estoque",
    "aging",
    "disponivel",
    "cotac",
    "orcament",
    "proposta",
)
TRUTHY = {"1", "true", "t", "yes", "y", "sim", "s"}
NUMERIC_COLUMN_HINTS = (
    "valor",
    "total",
    "qtd",
    "quant",
    "saldo",
    "estoque",
    "peso",
    "kg",
    "ton",
    "preco",
    "custo",
)
CATEGORY_COLUMN_HINTS = (
    "produto",
    "familia",
    "grupo",
    "vendedor",
    "cliente",
    "status",
    "situacao",
    "cidade",
    "uf",
    "deposito",
    "empresa",
)


def env_bool(env: dict[str, str], key: str, default: bool = False) -> bool:
    raw = env.get(key)
    if raw is None or raw == "":
        return default
    return raw.strip().lower() in TRUTHY


def json_default(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " / ").strip()


def normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value or "")
    return "".join(char for char in normalized if not unicodedata.combining(char)).lower().strip()


def parse_compact_number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if number == number and number not in (float("inf"), float("-inf")) else None
    text = str(value).strip()
    if not text:
        return None
    cleaned = "".join(char for char in text if char.isdigit() or char in ",.-")
    if not cleaned or cleaned in {"-", ".", ","}:
        return None
    if "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return None


def parse_spreadsheet_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    candidates = (
        (text[:10], "%Y-%m-%d"),
        (text[:10], "%d/%m/%Y"),
        (text[:10], "%d-%m-%Y"),
        (text[:19], "%Y-%m-%d %H:%M:%S"),
        (text[:19], "%d/%m/%Y %H:%M:%S"),
    )
    for candidate, fmt in candidates:
        try:
            return datetime.strptime(candidate, fmt).date()
        except ValueError:
            continue
    return None


def normalized_lookup(payload: dict[str, Any], aliases: tuple[str, ...]) -> Any:
    normalized_aliases = {normalize_text(alias) for alias in aliases}
    normalized_payload = {normalize_text(key): value for key, value in payload.items()}
    for alias in normalized_aliases:
        if alias in normalized_payload:
            return normalized_payload[alias]
    for key, value in normalized_payload.items():
        if any(alias in key or key in alias for alias in normalized_aliases):
            return value
    return None


def clean_text_value(value: Any) -> str | None:
    text = clean_cell(value)
    return text or None


def payload_from_staging(row: dict[str, Any]) -> dict[str, Any]:
    payload = row.get("dados_transformados")
    if isinstance(payload, dict):
        return payload
    envelope = row.get("payload_original")
    if isinstance(envelope, dict) and isinstance(envelope.get("data"), dict):
        return envelope["data"]
    return {}


def file_key(file_info: dict[str, Any]) -> str:
    modified = file_info.get("modifiedTime") or "sem_data"
    return f"{file_info['id']}:{modified}"


def is_supported_spreadsheet(file_info: dict[str, Any]) -> bool:
    mime_type = file_info.get("mimeType") or ""
    name = file_info.get("name") or ""
    lowered = normalize_text(name)
    if any(lowered.endswith(suffix) for suffix in GENERATED_ARCHIVE_SUFFIXES):
        return False
    if lowered.startswith("_") or "nao excluir" in lowered:
        return False
    if not any(keyword in lowered for keyword in BUSINESS_SPREADSHEET_KEYWORDS):
        return False
    suffix = Path(name).suffix.lower()
    return mime_type == GOOGLE_SHEETS_MIME or mime_type == XLSX_MIME or mime_type in CSV_MIME_TYPES or suffix in SUPPORTED_EXTENSIONS


def list_spreadsheets_recursive(token: str, root_folder_id: str, *, max_depth: int = 5) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    queue: list[tuple[str, int]] = [(root_folder_id, 0)]
    seen_folders: set[str] = set()
    while queue:
        folder_id, depth = queue.pop(0)
        if folder_id in seen_folders or depth > max_depth:
            continue
        seen_folders.add(folder_id)
        for child in list_drive_children(token, folder_id):
            mime_type = child.get("mimeType")
            if mime_type == DRIVE_FOLDER_MIME:
                queue.append((child["id"], depth + 1))
                continue
            if is_supported_spreadsheet(child):
                child["parent_id"] = folder_id
                found.append(child)
    return found


def configured_source_folders(env: dict[str, str], fallback_folder_id: str) -> list[dict[str, str]]:
    raw_ids = [
        env.get("DRIVE_SPREADSHEET_SOURCE_FOLDER_ID", "").strip(),
        env.get("GOOGLE_DRIVE_SPREADSHEET_FOLDER_ID", "").strip(),
    ]
    raw_urls = [
        env.get("DRIVE_SPREADSHEET_SOURCE_FOLDER_URL", "").strip(),
        env.get("GOOGLE_DRIVE_SPREADSHEET_FOLDER_URL", "").strip(),
    ]
    folders: list[dict[str, str]] = []
    for folder_id, folder_url in zip(raw_ids, raw_urls):
        if folder_id:
            folders.append({"id": folder_id, "url": folder_url or f"https://drive.google.com/drive/folders/{folder_id}", "role": "source"})
    if not folders and DEFAULT_SOURCE_FOLDER_ID:
        folders.append({"id": DEFAULT_SOURCE_FOLDER_ID, "url": DEFAULT_SOURCE_FOLDER_URL, "role": "source_default"})
    if not folders and fallback_folder_id:
        folders.append({"id": fallback_folder_id, "url": env.get("ARCHIVE_GOOGLE_DRIVE_FOLDER_URL", ""), "role": "archive_fallback"})
    deduped: list[dict[str, str]] = []
    seen: set[str] = set()
    for folder in folders:
        if folder["id"] in seen:
            continue
        seen.add(folder["id"])
        deduped.append(folder)
    return deduped


def already_processed(cur: Any, file_info: dict[str, Any]) -> bool:
    cur.execute(
        """
        select i.id
        from public.drive_spreadsheet_ingestions i
        join public.dashboard_drive_spreadsheet_cache c
          on c.drive_file_id = i.drive_file_id
         and c.drive_modified_time = i.drive_modified_time
        where i.drive_file_id = %s
          and i.drive_modified_time = %s::timestamptz
          and i.status = 'sucesso'
          and coalesce((c.payload->>'compact_schema_version')::int, 0) >= 3
        limit 1
        """,
        (file_info["id"], file_info.get("modifiedTime")),
    )
    return cur.fetchone() is not None


def start_ingestion(cur: Any, file_info: dict[str, Any], sync_id: str, local_path: Path | None = None) -> str:
    cur.execute(
        """
        insert into public.drive_spreadsheet_ingestions(
          drive_file_id, drive_file_name, drive_mime_type, drive_modified_time,
          drive_parent_id, local_path, sync_id, metadata
        )
        values (%s, %s, %s, %s::timestamptz, %s, %s, %s, %s::jsonb)
        on conflict (drive_file_id, drive_modified_time)
        do update set
          status = 'processando',
          sync_id = excluded.sync_id,
          started_at = now(),
          finished_at = null,
          error_message = null,
          local_path = excluded.local_path
        returning id::text
        """,
        (
            file_info["id"],
            file_info.get("name"),
            file_info.get("mimeType"),
            file_info.get("modifiedTime"),
            file_info.get("parent_id"),
            str(local_path.relative_to(ROOT)) if local_path else None,
            sync_id,
            Json({"webViewLink": file_info.get("webViewLink"), "size": file_info.get("size")}),
        ),
    )
    return str(cur.fetchone()[0])


def finish_ingestion(
    cur: Any,
    ingestion_id: str,
    *,
    status: str,
    rows_read: int = 0,
    rows_inserted: int = 0,
    worksheets: list[dict[str, Any]] | None = None,
    file_hash: str | None = None,
    detected_type: str = "GERAL",
    error_message: str | None = None,
) -> None:
    cur.execute(
        """
        update public.drive_spreadsheet_ingestions
        set status = %s,
            rows_read = %s,
            rows_inserted = %s,
            worksheets = %s::jsonb,
            file_sha256 = %s,
            detected_type = %s,
            error_message = %s,
            finished_at = now()
        where id = %s::uuid
        """,
        (
            status,
            rows_read,
            rows_inserted,
            Json(worksheets or []),
            file_hash,
            detected_type,
            error_message,
            ingestion_id,
        ),
    )


def build_compact_summary(
    *,
    file_info: dict[str, Any],
    detected_type: str,
    worksheets: list[dict[str, Any]],
    staging_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    header_counter: Counter[str] = Counter()
    numeric_totals: dict[str, dict[str, float | int]] = {}
    category_counts: dict[str, Counter[str]] = {}
    sample_rows: list[dict[str, Any]] = []

    for staging_row in staging_rows:
        envelope = staging_row.get("payload_original") or {}
        if not isinstance(envelope, dict):
            continue
        payload = envelope.get("data")
        if not isinstance(payload, dict):
            payload = envelope
        if len(sample_rows) < 3:
            sample_rows.append({key: payload.get(key) for key in list(payload.keys())[:12]})
        for key, value in payload.items():
            if key in {"_drive_file_id", "_drive_file_name", "_sheet_name", "_row_number", "_detected_type"}:
                continue
            header_counter[key] += 1
            normalized_key = normalize_text(key)
            if any(hint in normalized_key for hint in NUMERIC_COLUMN_HINTS):
                number = parse_compact_number(value)
                if number is not None:
                    stats = numeric_totals.setdefault(key, {"sum": 0.0, "count": 0})
                    stats["sum"] = float(stats["sum"]) + number
                    stats["count"] = int(stats["count"]) + 1
            if any(hint in normalized_key for hint in CATEGORY_COLUMN_HINTS):
                text = str(value or "").strip()
                if text:
                    category_counts.setdefault(key, Counter())[text[:120]] += 1

    numeric_rows = [
        {"campo": key, "soma": round(float(value["sum"]), 2), "preenchidos": int(value["count"])}
        for key, value in numeric_totals.items()
        if int(value["count"]) > 0
    ]
    numeric_rows.sort(key=lambda item: abs(float(item["soma"])), reverse=True)

    category_rows = []
    for key, counter in category_counts.items():
        if not counter:
            continue
        category_rows.append(
            {
                "campo": key,
                "valores": [{"valor": value, "linhas": count} for value, count in counter.most_common(8)],
            }
        )
    category_rows.sort(key=lambda item: sum(value["linhas"] for value in item["valores"]), reverse=True)

    return {
        "detected_type": detected_type,
        "drive_file_id": file_info["id"],
        "drive_file_name": file_info.get("name"),
        "drive_modified_time": file_info.get("modifiedTime"),
        "source_folder_id": file_info.get("source_folder_id") or file_info.get("parent_id"),
        "source_folder_url": file_info.get("source_folder_url"),
        "rows": len(staging_rows),
        "worksheets": worksheets,
        "headers": [{"campo": key, "linhas": count} for key, count in header_counter.most_common(30)],
        "numeric_totals": numeric_rows[:16],
        "top_categories": category_rows[:8],
        "sample_rows": sample_rows,
        "compact_schema_version": 3,
    }


def upsert_dashboard_cache(cur: Any, file_info: dict[str, Any], detected_type: str, payload: dict[str, Any]) -> None:
    cache_key = f"drive_spreadsheet:{detected_type}:{file_info['id']}:{file_info.get('modifiedTime') or 'sem_data'}"
    cur.execute(
        """
        insert into public.dashboard_drive_spreadsheet_cache(
          cache_key, detected_type, drive_file_id, drive_file_name,
          drive_modified_time, row_count, payload, refreshed_at
        )
        values (%s, %s, %s, %s, %s::timestamptz, %s, %s::jsonb, now())
        on conflict (cache_key)
        do update set
          detected_type = excluded.detected_type,
          drive_file_name = excluded.drive_file_name,
          drive_modified_time = excluded.drive_modified_time,
          row_count = excluded.row_count,
          payload = excluded.payload,
          refreshed_at = now()
        """,
        (
            cache_key,
            detected_type,
            file_info["id"],
            file_info.get("name") or file_info["id"],
            file_info.get("modifiedTime"),
            int(payload.get("rows") or 0),
            Json(payload),
        ),
    )


def detect_type(file_name: str, sheets: list[str]) -> str:
    probe = normalize_text(f"{file_name} {' '.join(sheets)}")
    if "envelhecimento" in probe or "aging" in probe or ("env" in probe and "estoque" in probe):
        return "ENVELHECIMENTO_ESTOQUE"
    if "estoque" in probe and ("disponivel" in probe or "saldo" in probe):
        return "ESTOQUE_DISPONIVEL"
    if "cotac" in probe or "orcament" in probe or "proposta" in probe:
        return "COTACAO"
    if "gest" in probe and ("produ" in probe or "prod" in probe):
        return "GESTAO_PRODUCAO"
    if "margem" in probe or "gabr" in probe:
        return "MARGEM"
    if "ranking" in probe:
        return "RANKING"
    if "estoque" in probe:
        return "ESTOQUE"
    if "cliente" in probe:
        return "CLIENTES"
    return "GERAL"


def infer_header_row(file_name: str, sheet_name: str, preview_rows: list[tuple[Any, ...]]) -> int:
    configured = HEADER_ROWS.get(file_name, {}).get(sheet_name)
    if configured:
        return configured
    best_index = 1
    best_score = -1
    for index, row in enumerate(preview_rows[:12], start=1):
        headers = [clean_cell(value) for value in row]
        filled = [header for header in headers if header]
        score = len(filled)
        if any(keyword in " ".join(filled).lower() for keyword in ("cliente", "produto", "pedido", "data", "valor", "nota", "vendedor")):
            score += 10
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def row_hash(payload: dict[str, Any]) -> str:
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True, default=json_default).encode("utf-8")).hexdigest()


def worksheet_rows(path: Path, *, drive_file_id: str, drive_file_name: str, sync_id: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]], str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheets: list[dict[str, Any]] = []
    staging_rows: list[dict[str, Any]] = []
    sheet_names = [sheet.title for sheet in workbook.worksheets]
    detected = detect_type(drive_file_name, sheet_names)

    for worksheet in workbook.worksheets:
        max_row = worksheet.max_row or 0
        if max_row < 1:
            continue
        preview_rows = list(worksheet.iter_rows(min_row=1, max_row=min(12, max_row), values_only=True))
        header_row = infer_header_row(drive_file_name, worksheet.title, preview_rows)
        header_values = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        headers = [clean_cell(value) or f"coluna_{index + 1}" for index, value in enumerate(header_values)]
        rows_in_sheet = 0
        rows_insertable = 0

        for row_number, values in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not values or not any(value not in (None, "") for value in values):
                continue
            rows_in_sheet += 1
            payload = {
                headers[index] if index < len(headers) else f"coluna_{index + 1}": clean_cell(value)
                for index, value in enumerate(values)
            }
            payload = {key: value for key, value in payload.items() if key and value != ""}
            if not payload:
                continue
            rows_insertable += 1
            staging_payload = {
                "drive_file_id": drive_file_id,
                "drive_file_name": drive_file_name,
                "sheet_name": worksheet.title,
                "row_number": row_number,
                "tipo_planilha": detected,
                "data": payload,
            }
            staging_rows.append(
                {
                    "source_system": "GOOGLE_DRIVE_ARCHIVE",
                    "source_id": f"{drive_file_id}:{worksheet.title}:{row_number}",
                    "entidade": f"planilha_{detected.lower()}",
                    "payload_original": staging_payload,
                    "dados_transformados": payload,
                    "sync_id": sync_id,
                    "hash_registro": row_hash(staging_payload),
                    "tabela_destino": None,
                }
            )
        worksheets.append(
            {
                "name": worksheet.title,
                "rows": worksheet.max_row,
                "cols": worksheet.max_column,
                "header_row": header_row,
                "rows_read": rows_in_sheet,
                "rows_insertable": rows_insertable,
                "headers": headers,
            }
        )
    workbook.close()
    return staging_rows, worksheets, detected


def csv_rows(path: Path, *, drive_file_id: str, drive_file_name: str, sync_id: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]], str]:
    import csv

    detected = detect_type(drive_file_name, [path.stem])
    staging_rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as file_obj:
        reader = csv.DictReader(file_obj)
        headers = [clean_cell(header) for header in (reader.fieldnames or [])]
        for row_number, row in enumerate(reader, start=2):
            payload = {clean_cell(key): clean_cell(value) for key, value in row.items() if clean_cell(key) and clean_cell(value) != ""}
            if not payload:
                continue
            staging_payload = {
                "drive_file_id": drive_file_id,
                "drive_file_name": drive_file_name,
                "sheet_name": path.stem,
                "row_number": row_number,
                "tipo_planilha": detected,
                "data": payload,
            }
            staging_rows.append(
                {
                    "source_system": "GOOGLE_DRIVE_ARCHIVE",
                    "source_id": f"{drive_file_id}:{path.stem}:{row_number}",
                    "entidade": f"planilha_{detected.lower()}",
                    "payload_original": staging_payload,
                    "dados_transformados": payload,
                    "sync_id": sync_id,
                    "hash_registro": row_hash(staging_payload),
                    "tabela_destino": None,
                }
            )
    worksheets = [
        {
            "name": path.stem,
            "rows": len(staging_rows) + 1,
            "cols": len(headers),
            "header_row": 1,
            "rows_read": len(staging_rows),
            "rows_insertable": len(staging_rows),
            "headers": headers,
        }
    ]
    return staging_rows, worksheets, detected


def insert_staging_rows(cur: Any, rows: list[dict[str, Any]]) -> int:
    if not rows:
        return 0
    inserted = 0
    cur.execute(
        """
        create temporary table if not exists tmp_drive_spreadsheet_rows (
          source_system text,
          source_id text,
          entidade text,
          payload_original jsonb,
          dados_transformados jsonb,
          sync_id text,
          hash_registro text,
          tabela_destino text
        ) on commit drop
        """
    )
    cur.execute("truncate table tmp_drive_spreadsheet_rows")
    batch_size = 5000
    for start in range(0, len(rows), batch_size):
        batch = rows[start : start + batch_size]
        cur.executemany(
            """
            insert into tmp_drive_spreadsheet_rows(
              source_system, source_id, entidade, payload_original, dados_transformados,
              sync_id, hash_registro, tabela_destino
            )
            values (%s, %s, %s, %s::jsonb, %s::jsonb, %s, %s, %s)
            """,
            [
                (
                    row["source_system"],
                    row["source_id"],
                    row["entidade"],
                    Json(row["payload_original"]),
                    Json(row["dados_transformados"]),
                    row["sync_id"],
                    row["hash_registro"],
                    row["tabela_destino"],
                )
                for row in batch
            ],
        )
        cur.execute(
            """
            insert into public.staging_dados(
              source_system, source_id, entidade, payload_original, dados_transformados,
              sync_id, hash_registro, status_validacao, tabela_destino
            )
            select
              t.source_system,
              t.source_id,
              t.entidade,
              t.payload_original,
              t.dados_transformados,
              t.sync_id,
              t.hash_registro,
              'valido',
              t.tabela_destino
            from tmp_drive_spreadsheet_rows t
            """
        )
        inserted += int(cur.rowcount or 0)
        cur.execute("truncate table tmp_drive_spreadsheet_rows")
    return inserted


def replace_file_fact_rows(cur: Any, table_name: str, file_info: dict[str, Any]) -> None:
    cur.execute(
        f"""
        delete from public.{table_name}
        where drive_file_id = %s
          and drive_modified_time = %s::timestamptz
        """,
        (file_info["id"], file_info.get("modifiedTime")),
    )


def insert_cotacao_facts(cur: Any, file_info: dict[str, Any], rows: list[dict[str, Any]], sync_id: str) -> int:
    fact_rows = []
    for row in rows:
        payload = payload_from_staging(row)
        if not payload:
            continue
        fact_rows.append(
            (
                row["source_id"],
                file_info["id"],
                file_info.get("name"),
                file_info.get("modifiedTime"),
                sync_id,
                normalized_lookup(row.get("payload_original", {}), ("row_number", "_row_number")),
                parse_spreadsheet_date(normalized_lookup(payload, ("DataCotacao", "Data Cotacao", "Data da Cotacao"))),
                clean_text_value(normalized_lookup(payload, ("NumeroCotacao", "Numero Cotacao", "Cotacao", "Orcamento"))),
                clean_text_value(normalized_lookup(payload, ("Numero Esboco", "NumeroEsboco", "Esboco"))),
                clean_text_value(normalized_lookup(payload, ("Chave Esboco", "ChaveEsboco"))),
                clean_text_value(normalized_lookup(payload, ("Pedido Destino", "PedidoDestino", "Pedido"))),
                parse_spreadsheet_date(normalized_lookup(payload, ("DataAdicaoPV", "Data Adicao PV", "Data Pedido"))),
                parse_spreadsheet_date(normalized_lookup(payload, ("DataNF", "Data NF", "Data Nota"))),
                clean_text_value(normalized_lookup(payload, ("NumeroNF", "Numero NF", "NF"))),
                clean_text_value(normalized_lookup(payload, ("Status", "Situacao"))),
                clean_text_value(normalized_lookup(payload, ("Vendedor", "Representante"))),
                clean_text_value(normalized_lookup(payload, ("Unidade", "Empresa"))),
                clean_text_value(normalized_lookup(payload, ("CodCliente", "Codigo Cliente", "Cod. Cliente"))),
                clean_text_value(normalized_lookup(payload, ("Cliente", "Nome Cliente"))),
                clean_text_value(normalized_lookup(payload, ("Cidade", "Municipio"))),
                clean_text_value(normalized_lookup(payload, ("UF", "Estado"))),
                clean_text_value(normalized_lookup(payload, ("TipoFrete", "Tipo Frete", "Frete"))),
                clean_text_value(normalized_lookup(payload, ("Item_PA", "Item PA", "Codigo Produto", "Item"))),
                clean_text_value(normalized_lookup(payload, ("Desc_PA", "Desc PA", "Descricao Produto", "Produto", "Descricao"))),
                clean_text_value(normalized_lookup(payload, ("Familia", "Família"))),
                parse_compact_number(normalized_lookup(payload, ("QtdPedido", "Qtd Pedido", "Quantidade"))),
                parse_compact_number(normalized_lookup(payload, ("Peso", "Peso Kg", "Peso KG", "Kg"))),
                parse_compact_number(normalized_lookup(payload, ("PrecoKg", "Preco Kg", "Preço Kg", "R$/Kg"))),
                parse_compact_number(normalized_lookup(payload, ("Total", "Valor Total", "Valor"))),
                parse_spreadsheet_date(normalized_lookup(payload, ("DataEntregaComercial", "Data Entrega Comercial", "Entrega"))),
                parse_compact_number(normalized_lookup(payload, ("Estoque_PA", "Estoque PA", "Estoque"))),
                parse_compact_number(normalized_lookup(payload, ("EstoqueConfirmado_PA", "Estoque Confirmado PA", "Estoque Confirmado"))),
                Json(payload),
            )
        )
    if not fact_rows:
        return 0
    replace_file_fact_rows(cur, "fato_cotacao_item", file_info)
    cur.executemany(
        """
        insert into public.fato_cotacao_item(
          source_id, drive_file_id, drive_file_name, drive_modified_time, sync_id, row_number,
          data_cotacao, numero_cotacao, numero_esboco, chave_esboco, pedido_destino,
          data_adicao_pv, data_nf, numero_nf, status_original, vendedor, unidade,
          cod_cliente, cliente, cidade, uf, tipo_frete, item_pa, desc_pa, familia,
          qtd_pedido, peso_kg, preco_kg, valor_total, data_entrega_comercial,
          estoque_pa, estoque_confirmado_pa, payload_original
        )
        values (
          %s, %s, %s, %s::timestamptz, %s, %s,
          %s, %s, %s, %s, %s,
          %s, %s, %s, %s, %s, %s,
          %s, %s, %s, %s, %s, %s, %s, %s,
          %s, %s, %s, %s, %s,
          %s, %s, %s::jsonb
        )
        on conflict (source_id) do update set
          drive_file_name = excluded.drive_file_name,
          drive_modified_time = excluded.drive_modified_time,
          sync_id = excluded.sync_id,
          data_cotacao = excluded.data_cotacao,
          status_original = excluded.status_original,
          valor_total = excluded.valor_total,
          peso_kg = excluded.peso_kg,
          payload_original = excluded.payload_original
        """,
        fact_rows,
    )
    return len(fact_rows)


def insert_estoque_disponivel_facts(cur: Any, file_info: dict[str, Any], rows: list[dict[str, Any]], sync_id: str) -> int:
    fact_rows = []
    for row in rows:
        payload = payload_from_staging(row)
        if not payload:
            continue
        fact_rows.append(
            (
                row["source_id"],
                file_info["id"],
                file_info.get("name"),
                file_info.get("modifiedTime"),
                sync_id,
                normalized_lookup(row.get("payload_original", {}), ("row_number", "_row_number")),
                clean_text_value(normalized_lookup(payload, ("Familia", "Família"))),
                clean_text_value(normalized_lookup(payload, ("SubGrupo", "Subgrupo", "Grupo"))),
                clean_text_value(normalized_lookup(payload, ("Codigo", "Código", "Cod. material", "Cod Material"))),
                clean_text_value(normalized_lookup(payload, ("Descricao", "Descrição", "Desc. material", "Desc Material"))),
                clean_text_value(normalized_lookup(payload, ("Laminação", "Laminacao"))),
                clean_text_value(normalized_lookup(payload, ("Espessura",))),
                parse_compact_number(normalized_lookup(payload, ("Estoque Disponivel (Kg)", "Estoque Disponivel Kg", "Estoque Disponível (Kg)", "Estoque"))),
                clean_text_value(normalized_lookup(payload, ("Deposito", "Depósito"))),
                clean_text_value(normalized_lookup(payload, ("NomeDeposito", "Nome Deposito", "Descrição - Deposito"))),
                Json(payload),
            )
        )
    if not fact_rows:
        return 0
    replace_file_fact_rows(cur, "fato_estoque_disponivel", file_info)
    cur.executemany(
        """
        insert into public.fato_estoque_disponivel(
          source_id, drive_file_id, drive_file_name, drive_modified_time, sync_id, row_number,
          familia, subgrupo, codigo, descricao, laminacao, espessura, estoque_disponivel_kg,
          deposito, nome_deposito, payload_original
        )
        values (%s, %s, %s, %s::timestamptz, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
        on conflict (source_id) do update set
          drive_file_name = excluded.drive_file_name,
          drive_modified_time = excluded.drive_modified_time,
          sync_id = excluded.sync_id,
          estoque_disponivel_kg = excluded.estoque_disponivel_kg,
          payload_original = excluded.payload_original
        """,
        fact_rows,
    )
    return len(fact_rows)


def insert_estoque_envelhecimento_facts(cur: Any, file_info: dict[str, Any], rows: list[dict[str, Any]], sync_id: str) -> int:
    fact_rows = []
    for row in rows:
        payload = payload_from_staging(row)
        if not payload:
            continue
        fact_rows.append(
            (
                row["source_id"],
                file_info["id"],
                file_info.get("name"),
                file_info.get("modifiedTime"),
                sync_id,
                normalized_lookup(row.get("payload_original", {}), ("row_number", "_row_number")),
                parse_spreadsheet_date(normalized_lookup(payload, ("Dt. Base", "Data Base", "Dt Base"))),
                clean_text_value(normalized_lookup(payload, ("Unidade", "Empresa"))),
                clean_text_value(normalized_lookup(payload, ("Cod. material", "Cod Material", "Codigo"))),
                clean_text_value(normalized_lookup(payload, ("Familia", "Família"))),
                clean_text_value(normalized_lookup(payload, ("Desc. material", "Desc Material", "Descricao"))),
                clean_text_value(normalized_lookup(payload, ("Depósito", "Deposito"))),
                clean_text_value(normalized_lookup(payload, ("Descrição - Deposito", "Descricao Deposito", "NomeDeposito"))),
                clean_text_value(normalized_lookup(payload, ("Aging", "Faixa Aging", "Idade"))),
                parse_compact_number(normalized_lookup(payload, ("Quantidade", "Quantidade Kg", "Qtd", "Kg"))),
                parse_compact_number(normalized_lookup(payload, ("Valor", "Valor Total"))),
                parse_compact_number(normalized_lookup(payload, ("Custo médio por kg", "Custo medio por kg", "Custo Medio Kg"))),
                clean_text_value(normalized_lookup(payload, ("Grupo",))),
                clean_text_value(normalized_lookup(payload, ("Tipo",))),
                clean_text_value(normalized_lookup(payload, ("Mercado",))),
                Json(payload),
            )
        )
    if not fact_rows:
        return 0
    replace_file_fact_rows(cur, "fato_estoque_envelhecimento", file_info)
    cur.executemany(
        """
        insert into public.fato_estoque_envelhecimento(
          source_id, drive_file_id, drive_file_name, drive_modified_time, sync_id, row_number,
          dt_base, unidade, cod_material, familia, desc_material, deposito, descricao_deposito,
          aging, quantidade_kg, valor, custo_medio_kg, grupo, tipo, mercado, payload_original
        )
        values (%s, %s, %s, %s::timestamptz, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
        on conflict (source_id) do update set
          drive_file_name = excluded.drive_file_name,
          drive_modified_time = excluded.drive_modified_time,
          sync_id = excluded.sync_id,
          quantidade_kg = excluded.quantidade_kg,
          valor = excluded.valor,
          payload_original = excluded.payload_original
        """,
        fact_rows,
    )
    return len(fact_rows)


def insert_business_facts(cur: Any, detected_type: str, file_info: dict[str, Any], rows: list[dict[str, Any]], sync_id: str) -> int:
    if detected_type == "COTACAO":
        return insert_cotacao_facts(cur, file_info, rows, sync_id)
    if detected_type == "ESTOQUE_DISPONIVEL":
        return insert_estoque_disponivel_facts(cur, file_info, rows, sync_id)
    if detected_type == "ENVELHECIMENTO_ESTOQUE":
        return insert_estoque_envelhecimento_facts(cur, file_info, rows, sync_id)
    return 0


def download_spreadsheet(token: str, file_info: dict[str, Any], target_dir: Path) -> Path:
    name = file_info.get("name") or file_info["id"]
    mime_type = file_info.get("mimeType")
    suffix = Path(name).suffix.lower()
    base_name = slug(Path(name).stem or file_info["id"])
    modified = (file_info.get("modifiedTime") or "sem_data").replace(":", "").replace("-", "")
    if mime_type == GOOGLE_SHEETS_MIME:
        destination = target_dir / f"{base_name}_{modified}.xlsx"
        export_drive_file(token, file_info["id"], destination, XLSX_MIME)
        return destination
    if suffix not in SUPPORTED_EXTENSIONS:
        destination = target_dir / f"{base_name}_{modified}.xlsx"
    else:
        destination = target_dir / f"{base_name}_{modified}{suffix}"
    download_drive_file(token, file_info["id"], destination)
    return destination


def collect_drive_spreadsheets(
    *,
    execute: bool = True,
    max_files: int | None = None,
    max_depth: int = DEFAULT_SCAN_DEPTH,
    force: bool = False,
) -> dict[str, Any]:
    env = load_env()
    config = archive_config(env)
    if config.provider != "google_drive" or not config.drive_folder_id:
        return {"status": "skipped", "reason": "Archive Google Drive nao configurado."}
    token = drive_access_token(config)
    if not token:
        raise RuntimeError("Credencial Google OAuth/Service Account nao configurada para Drive.")

    sync_id = f"drive_spreadsheets_{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')}"
    store_raw_rows = env_bool(env, "DRIVE_SPREADSHEET_STORE_RAW_ROWS", False)
    source_folders = configured_source_folders(env, config.drive_folder_id)
    files: list[dict[str, Any]] = []
    for folder in source_folders:
        folder_files = list_spreadsheets_recursive(token, folder["id"], max_depth=max_depth)
        for file_info in folder_files:
            file_info["source_folder_id"] = folder["id"]
            file_info["source_folder_url"] = folder["url"]
            file_info["source_folder_role"] = folder["role"]
        files.extend(folder_files)
    files = sorted(files, key=lambda item: item.get("modifiedTime") or "", reverse=True)
    if max_files is not None:
        files = files[:max_files]

    target_dir = (ROOT / config.local_dir / "google_drive_archive" / "incoming_spreadsheets").resolve()
    processed: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    with connect_core_database(env) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                insert into public.fontes_dados(nome, tipo, classificacao, status, frequencia_sincronizacao, url, descricao)
                values ('Google Drive Planilhas Brutas', 'csv_upload', 'secundaria', 'ativa', 'diaria', %s, 'Pasta de entrada monitorada automaticamente para planilhas brutas.')
                on conflict (nome) do update set
                  status = 'ativa',
                  frequencia_sincronizacao = 'diaria',
                  url = excluded.url,
                  atualizado_em = now()
                """,
                (source_folders[0]["url"] if source_folders else env.get("ARCHIVE_GOOGLE_DRIVE_FOLDER_URL") or config.drive_folder_id,),
            )
            for file_info in files:
                if not force and already_processed(cur, file_info):
                    skipped.append({"file_id": file_info["id"], "name": file_info.get("name"), "reason": "unchanged"})
                    continue
                local_path: Path | None = None
                ingestion_id = start_ingestion(cur, file_info, sync_id)
                try:
                    local_path = download_spreadsheet(token, file_info, target_dir)
                    cur.execute(
                        "update public.drive_spreadsheet_ingestions set local_path = %s where id = %s::uuid",
                        (str(local_path.relative_to(ROOT)), ingestion_id),
                    )
                    if local_path.suffix.lower() == ".csv":
                        staging_rows, worksheets, detected = csv_rows(
                            local_path,
                            drive_file_id=file_info["id"],
                            drive_file_name=file_info.get("name") or file_info["id"],
                            sync_id=sync_id,
                        )
                    else:
                        staging_rows, worksheets, detected = worksheet_rows(
                            local_path,
                            drive_file_id=file_info["id"],
                            drive_file_name=file_info.get("name") or file_info["id"],
                            sync_id=sync_id,
                        )
                    compact_summary = build_compact_summary(
                        file_info=file_info,
                        detected_type=detected,
                        worksheets=worksheets,
                        staging_rows=staging_rows,
                    )
                    if execute:
                        upsert_dashboard_cache(cur, file_info, detected, compact_summary)
                    facts_inserted = insert_business_facts(cur, detected, file_info, staging_rows, sync_id) if execute else 0
                    inserted = insert_staging_rows(cur, staging_rows) if execute and store_raw_rows else 0
                    finish_ingestion(
                        cur,
                        ingestion_id,
                        status="sucesso",
                        rows_read=len(staging_rows),
                        rows_inserted=inserted,
                        worksheets=worksheets,
                        file_hash=file_sha256(local_path),
                        detected_type=detected,
                    )
                    processed.append(
                        {
                            "file_id": file_info["id"],
                            "name": file_info.get("name"),
                            "modified_time": file_info.get("modifiedTime"),
                            "detected_type": detected,
                            "rows_read": len(staging_rows),
                            "rows_inserted": inserted,
                            "facts_inserted": facts_inserted,
                            "raw_rows_stored": bool(store_raw_rows),
                            "local_path": str(local_path.relative_to(ROOT)),
                        }
                    )
                except Exception as exc:
                    message = f"{type(exc).__name__}: {str(exc)[:500]}"
                    finish_ingestion(cur, ingestion_id, status="erro", error_message=message)
                    errors.append({"file_id": file_info["id"], "name": file_info.get("name"), "error": message})
            cur.execute(
                """
                update public.fontes_dados
                set ultima_sincronizacao = now(),
                    total_registros = %s,
                    registros_erro = %s,
                    status = %s,
                    atualizado_em = now()
                where nome = 'Google Drive Planilhas Brutas'
                """,
                (
                    sum(item["rows_read"] for item in processed),
                    len(errors),
                    "erro" if errors and not processed else "ativa",
                ),
            )
        if execute:
            conn.commit()
        else:
            conn.rollback()

    return {
        "status": "sucesso" if not errors else "parcial",
        "sync_id": sync_id,
        "raw_rows_stored": bool(store_raw_rows),
        "source_folders": source_folders,
        "files_found": len(files),
        "processed": processed,
        "skipped": skipped,
        "errors": errors,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Monitora a pasta de entrada do Google Drive e ingere planilhas novas.")
    parser.add_argument("--dry-run", action="store_true", help="Baixa e le planilhas, mas nao insere linhas em staging_dados.")
    parser.add_argument("--max-files", type=int, default=None, help="Limita quantidade de arquivos para diagnostico.")
    parser.add_argument("--max-depth", type=int, default=DEFAULT_SCAN_DEPTH, help="Profundidade maxima de subpastas no Drive.")
    parser.add_argument("--force", action="store_true", help="Reprocessa arquivos ja vistos para reconstruir caches/fatos.")
    args = parser.parse_args()

    result = collect_drive_spreadsheets(execute=not args.dry_run, max_files=args.max_files, max_depth=args.max_depth, force=args.force)
    print(json.dumps(result, ensure_ascii=False, indent=2, default=json_default))
    if result.get("errors"):
        sys.exit(2)


if __name__ == "__main__":
    main()
